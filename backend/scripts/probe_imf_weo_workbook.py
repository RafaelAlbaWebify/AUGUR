"""Read-only inspection of IMF's officially published October 2026 WEO workbook."""
import io
import json
import httpx
from openpyxl import load_workbook

URL="https://data.imf.org/-/media/iData/External-Storage/Documents/0FA3DC49A04D446E9D69DBEB29979E1C/en/WEOOct2026all.xlsx"


def inspect_workbook(payload: bytes) -> dict:
    if not payload.startswith(b"PK"):
        raise ValueError("IMF download is not an XLSX archive")
    book=load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
    try:
        sheets=[]
        for sheet in book.worksheets:
            samples=[]
            for row in sheet.iter_rows(min_row=1,max_row=min(12,sheet.max_row),max_col=min(18,sheet.max_column),values_only=True):
                samples.append([str(value)[:90] if value is not None else None for value in row])
            sheets.append({"name":sheet.title,"max_rows":sheet.max_row,"max_columns":sheet.max_column,
                           "sample":samples})
        return {"status":"read_only_source_inspection","format":"xlsx","sheets":sheets}
    finally:
        book.close()


def main():
    with httpx.Client(timeout=120,follow_redirects=True) as client:
        r=client.get(URL,headers={"Accept":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","User-Agent":"AUGUR-OfficialDataProbe/1.0"})
        print("IMF official workbook HTTP:",r.status_code,"size:",len(r.content))
        r.raise_for_status()
        result=inspect_workbook(r.content)
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=="__main__":
    main()
